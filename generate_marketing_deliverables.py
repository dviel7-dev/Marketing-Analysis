from __future__ import annotations

from pathlib import Path
from typing import Dict, List
import json

import matplotlib.pyplot as plt
import nbformat as nbf
import pandas as pd
import plotly.express as px
import seaborn as sns
from docx import Document
from docx.shared import Inches
from pptx import Presentation
from pptx.util import Inches as PptInches
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Image as PdfImage
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


BASE_DIR = Path(__file__).resolve().parent
DATA_FILE = BASE_DIR / "Marketing_Campaign_Data.csv"
OUTPUT_DIR = BASE_DIR / "deliverables"
CHART_DIR = OUTPUT_DIR / "charts"


# Use a consistent visual style across all generated assets.
sns.set_theme(style="whitegrid", context="talk")
plt.rcParams["figure.figsize"] = (12, 7)


def load_data() -> pd.DataFrame:
    df = pd.read_csv(DATA_FILE)
    df = df.rename(
        columns={
            "Interaction ID": "interaction_id",
            "Campaign Type": "campaign",
            "Channel": "channel",
            "Customer Type": "customer_type",
            "Converted (1=yes, 0=no)": "converted",
            "Time on Site (seconds)": "time_on_site",
            "Sales ($)": "sales",
        }
    )
    return df


def compute_kpis(df: pd.DataFrame) -> Dict[str, float]:
    converted_sales = df.loc[df["converted"] == 1, "sales"]
    kpis = {
        "interactions": float(len(df)),
        "total_revenue": float(df["sales"].sum()),
        "conversion_rate": float(df["converted"].mean()),
        "avg_time_on_site": float(df["time_on_site"].mean()),
        "avg_sale_converted": float(converted_sales.mean()),
        "new_customer_share": float((df["customer_type"] == "New").mean()),
    }
    return kpis


def grouped_tables(df: pd.DataFrame) -> Dict[str, pd.DataFrame]:
    by_campaign = (
        df.groupby("campaign", as_index=False)
        .agg(
            interactions=("interaction_id", "count"),
            conversion_rate=("converted", "mean"),
            revenue=("sales", "sum"),
            avg_time_on_site=("time_on_site", "mean"),
        )
        .sort_values("revenue", ascending=False)
    )

    by_channel = (
        df.groupby("channel", as_index=False)
        .agg(
            interactions=("interaction_id", "count"),
            conversion_rate=("converted", "mean"),
            revenue=("sales", "sum"),
            avg_time_on_site=("time_on_site", "mean"),
        )
        .sort_values("revenue", ascending=False)
    )

    by_combo = (
        df.groupby(["campaign", "channel"], as_index=False)
        .agg(
            interactions=("interaction_id", "count"),
            conversion_rate=("converted", "mean"),
            revenue=("sales", "sum"),
        )
        .sort_values("revenue", ascending=False)
    )

    new_customer_combo = (
        df[df["customer_type"] == "New"]
        .groupby(["campaign", "channel"], as_index=False)
        .agg(
            interactions=("interaction_id", "count"),
            conversion_rate=("converted", "mean"),
            revenue=("sales", "sum"),
        )
        .sort_values("revenue", ascending=False)
    )

    return {
        "by_campaign": by_campaign,
        "by_channel": by_channel,
        "by_combo": by_combo,
        "new_customer_combo": new_customer_combo,
    }


def save_charts(df: pd.DataFrame, tables: Dict[str, pd.DataFrame]) -> Dict[str, Path]:
    CHART_DIR.mkdir(parents=True, exist_ok=True)
    chart_paths: Dict[str, Path] = {}

    # Revenue by campaign and channel.
    pivot_revenue = tables["by_combo"].pivot(index="channel", columns="campaign", values="revenue")
    ax = pivot_revenue.plot(kind="bar", colormap="Set2")
    ax.set_title("Revenue by Channel and Campaign")
    ax.set_ylabel("Revenue ($)")
    ax.set_xlabel("Channel")
    ax.legend(title="Campaign")
    plt.tight_layout()
    path = CHART_DIR / "revenue_by_channel_campaign.png"
    plt.savefig(path, dpi=200)
    plt.close()
    chart_paths["revenue_by_channel_campaign"] = path

    # Conversion heatmap.
    pivot_conv = (
        tables["by_combo"].pivot(index="channel", columns="campaign", values="conversion_rate") * 100
    )
    plt.figure()
    sns.heatmap(pivot_conv, annot=True, fmt=".2f", cmap="YlGnBu")
    plt.title("Conversion Rate (%) by Channel and Campaign")
    plt.xlabel("Campaign")
    plt.ylabel("Channel")
    plt.tight_layout()
    path = CHART_DIR / "conversion_heatmap.png"
    plt.savefig(path, dpi=200)
    plt.close()
    chart_paths["conversion_heatmap"] = path

    # Revenue by customer type.
    customer_revenue = (
        df.groupby(["customer_type", "campaign"], as_index=False)["sales"].sum().sort_values("sales")
    )
    plt.figure()
    sns.barplot(data=customer_revenue, x="customer_type", y="sales", hue="campaign", palette="Set1")
    plt.title("Revenue by Customer Type and Campaign")
    plt.xlabel("Customer Type")
    plt.ylabel("Revenue ($)")
    plt.tight_layout()
    path = CHART_DIR / "customer_revenue.png"
    plt.savefig(path, dpi=200)
    plt.close()
    chart_paths["customer_revenue"] = path

    # Sales distribution for converted users.
    converted_df = df[df["converted"] == 1]
    plt.figure()
    sns.boxplot(data=converted_df, x="campaign", y="sales", hue="channel", palette="pastel")
    plt.title("Sales Distribution (Converted Users)")
    plt.xlabel("Campaign")
    plt.ylabel("Sales ($)")
    plt.tight_layout()
    path = CHART_DIR / "sales_distribution.png"
    plt.savefig(path, dpi=200)
    plt.close()
    chart_paths["sales_distribution"] = path

    # Time on site by conversion.
    plt.figure()
    sns.violinplot(
        data=df,
        x="campaign",
        y="time_on_site",
        hue="converted",
        split=True,
        palette="coolwarm",
    )
    plt.title("Time on Site Distribution by Campaign and Conversion")
    plt.xlabel("Campaign")
    plt.ylabel("Time on Site (seconds)")
    plt.tight_layout()
    path = CHART_DIR / "time_on_site_violin.png"
    plt.savefig(path, dpi=200)
    plt.close()
    chart_paths["time_on_site_violin"] = path

    return chart_paths


def build_dashboard_html(df: pd.DataFrame, tables: Dict[str, pd.DataFrame], kpis: Dict[str, float]) -> Path:
    output_path = OUTPUT_DIR / "marketing_campaign_executive_dashboard.html"
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    by_channel = tables["by_channel"].copy()
    by_channel["conversion_rate"] = by_channel["conversion_rate"] * 100

    by_campaign = tables["by_campaign"].copy()
    by_campaign["conversion_rate"] = by_campaign["conversion_rate"] * 100

    by_combo = tables["by_combo"].copy()
    by_combo["combo"] = by_combo["campaign"] + " - " + by_combo["channel"]
    by_combo = by_combo.sort_values("revenue", ascending=False)

    fig_channel_rev = px.bar(
        by_channel,
        x="channel",
        y="revenue",
        color="channel",
        title="Revenue by Channel",
        text_auto=".2s",
    )
    fig_channel_rev.update_layout(showlegend=False)

    fig_campaign_conv = px.bar(
        by_campaign,
        x="campaign",
        y="conversion_rate",
        color="campaign",
        title="Conversion Rate (%) by Campaign",
        text_auto=".2f",
    )
    fig_campaign_conv.update_layout(showlegend=False, yaxis_title="Conversion rate (%)")

    fig_combo = px.bar(
        by_combo,
        x="combo",
        y="revenue",
        color="campaign",
        title="Revenue by Campaign + Channel Combination",
    )
    fig_combo.update_layout(xaxis_title="Combination", yaxis_title="Revenue ($)")

    converted = df[df["converted"] == 1]
    fig_sales_dist = px.box(
        converted,
        x="channel",
        y="sales",
        color="campaign",
        title="Ticket Distribution (Converted Interactions)",
    )

    card_html = f"""
    <div class=\"cards\">
        <div class=\"card\"><h3>Interacciones Totales</h3><p>{int(kpis['interactions']):,}</p></div>
        <div class=\"card\"><h3>Ingresos Totales</h3><p>${kpis['total_revenue']:,.2f}</p></div>
        <div class=\"card\"><h3>Tasa de Conversion</h3><p>{kpis['conversion_rate'] * 100:.2f}%</p></div>
        <div class=\"card\"><h3>Tiempo Promedio en Sitio</h3><p>{kpis['avg_time_on_site']:.1f}s</p></div>
        <div class=\"card\"><h3>Ticket Promedio (Convertidos)</h3><p>${kpis['avg_sale_converted']:.2f}</p></div>
        <div class=\"card\"><h3>Participacion de Clientes Nuevos</h3><p>{kpis['new_customer_share'] * 100:.2f}%</p></div>
    </div>
    """

    executive_takeaways = """
    <div class=\"takeaways\">
        <h2>Conclusiones Ejecutivas</h2>
        <ul>
            <li>La Campana B aporta el mayor ingreso total, principalmente por su mayor volumen de interacciones.</li>
            <li>Email es el canal con mayor ingreso, mientras Instagram muestra la mejor eficiencia de conversion.</li>
            <li>La mejor combinacion para captar ingreso de clientes nuevos es Campana B + Email.</li>
            <li>El ticket promedio es estable entre campanas, por lo que el crecimiento depende de escalar trafico calificado y optimizar conversion.</li>
        </ul>
    </div>
    """

    html = f"""
    <!DOCTYPE html>
    <html lang=\"en\">
    <head>
      <meta charset=\"UTF-8\" />
      <meta name=\"viewport\" content=\"width=device-width, initial-scale=1.0\" />
    <title>Dashboard Ejecutivo de Campanas de Marketing</title>
      <script src=\"https://cdn.plot.ly/plotly-2.35.2.min.js\"></script>
      <style>
        :root {{
          --bg: #f4f1ea;
          --ink: #1f2a37;
          --accent: #005f73;
          --accent-2: #ca6702;
          --card: #ffffff;
        }}
        body {{
          margin: 0;
          font-family: Georgia, 'Times New Roman', serif;
          color: var(--ink);
          background: radial-gradient(circle at top right, #d9ead3, transparent 40%), var(--bg);
        }}
        .wrap {{
          max-width: 1280px;
          margin: 0 auto;
          padding: 28px;
        }}
        h1 {{
          margin: 0 0 8px;
          color: var(--accent);
          font-size: 2.2rem;
        }}
        .subtitle {{
          margin: 0 0 20px;
          color: #4a5568;
        }}
        .cards {{
          display: grid;
          grid-template-columns: repeat(auto-fit, minmax(180px, 1fr));
          gap: 14px;
          margin-bottom: 26px;
        }}
        .card {{
          background: var(--card);
          border-radius: 14px;
          padding: 14px;
          box-shadow: 0 7px 22px rgba(31, 42, 55, 0.08);
        }}
        .card h3 {{
          margin: 0;
          font-size: 0.9rem;
          color: #4a5568;
        }}
        .card p {{
          margin: 8px 0 0;
          font-size: 1.3rem;
          color: var(--accent-2);
          font-weight: 700;
        }}
        .plot-grid {{
          display: grid;
          grid-template-columns: 1fr;
          gap: 18px;
        }}
        .plot {{
          background: var(--card);
          border-radius: 14px;
          padding: 8px;
          box-shadow: 0 7px 22px rgba(31, 42, 55, 0.08);
        }}
        .takeaways {{
          margin-top: 22px;
          background: #ffffff;
          border-left: 6px solid var(--accent);
          border-radius: 10px;
          padding: 16px;
        }}
        .takeaways h2 {{ margin-top: 0; }}
        .takeaways li {{ margin-bottom: 8px; }}
      </style>
    </head>
    <body>
      <div class=\"wrap\">
        <h1>Dashboard Ejecutivo de Campanas de Marketing</h1>
        <p class=\"subtitle\">Dataset: 100,000 interacciones en Campana A y B con canales Email, Instagram y Web Banner.</p>
        {card_html}
        <div class=\"plot-grid\">
          <div class=\"plot\" id=\"plot1\"></div>
          <div class=\"plot\" id=\"plot2\"></div>
          <div class=\"plot\" id=\"plot3\"></div>
          <div class=\"plot\" id=\"plot4\"></div>
        </div>
        {executive_takeaways}
      </div>
      <script>
        Plotly.newPlot('plot1', {fig_channel_rev.to_json()}['data'], {fig_channel_rev.to_json()}['layout'], {{responsive: true}});
        Plotly.newPlot('plot2', {fig_campaign_conv.to_json()}['data'], {fig_campaign_conv.to_json()}['layout'], {{responsive: true}});
        Plotly.newPlot('plot3', {fig_combo.to_json()}['data'], {fig_combo.to_json()}['layout'], {{responsive: true}});
        Plotly.newPlot('plot4', {fig_sales_dist.to_json()}['data'], {fig_sales_dist.to_json()}['layout'], {{responsive: true}});
      </script>
    </body>
    </html>
    """

    output_path.write_text(html, encoding="utf-8")
    return output_path


def build_interactive_filtered_dashboard(df: pd.DataFrame) -> Path:
        output_path = OUTPUT_DIR / "marketing_campaign_dashboard_interactivo_filtros.html"
        OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

        filtered_df = df[["campaign", "channel", "customer_type", "converted", "time_on_site", "sales"]].copy()
        records_json = json.dumps(filtered_df.to_dict(orient="records"))

        html = f"""
        <!DOCTYPE html>
        <html lang=\"es\">
        <head>
            <meta charset=\"UTF-8\" />
            <meta name=\"viewport\" content=\"width=device-width, initial-scale=1.0\" />
            <title>Dashboard Interactivo con Filtros</title>
            <script src=\"https://cdn.plot.ly/plotly-2.35.2.min.js\"></script>
            <style>
                :root {{
                    --bg: #f9f7f1;
                    --card: #ffffff;
                    --ink: #1d3557;
                    --accent: #457b9d;
                    --accent2: #e76f51;
                }}
                body {{
                    margin: 0;
                    font-family: Cambria, Georgia, serif;
                    background: linear-gradient(135deg, #f6efe0 0%, #eef5f9 100%);
                    color: var(--ink);
                }}
                .container {{
                    max-width: 1280px;
                    margin: 0 auto;
                    padding: 24px;
                }}
                h1 {{
                    margin: 0 0 4px;
                    color: var(--ink);
                }}
                .subtitle {{
                    margin: 0 0 20px;
                    color: #4b5563;
                }}
                .filters {{
                    display: grid;
                    grid-template-columns: repeat(auto-fit, minmax(220px, 1fr));
                    gap: 12px;
                    margin-bottom: 16px;
                }}
                .filter-box {{
                    background: var(--card);
                    border-radius: 12px;
                    padding: 12px;
                    box-shadow: 0 6px 18px rgba(0, 0, 0, 0.08);
                }}
                label {{
                    display: block;
                    font-size: 0.9rem;
                    margin-bottom: 6px;
                }}
                select {{
                    width: 100%;
                    padding: 10px;
                    border-radius: 10px;
                    border: 1px solid #d1d5db;
                    font-size: 0.95rem;
                }}
                .cards {{
                    display: grid;
                    grid-template-columns: repeat(auto-fit, minmax(180px, 1fr));
                    gap: 10px;
                    margin-bottom: 14px;
                }}
                .card {{
                    background: var(--card);
                    border-radius: 12px;
                    padding: 12px;
                    box-shadow: 0 6px 18px rgba(0, 0, 0, 0.08);
                }}
                .card h3 {{
                    margin: 0;
                    font-size: 0.88rem;
                    color: #4b5563;
                }}
                .card p {{
                    margin: 8px 0 0;
                    font-size: 1.2rem;
                    color: var(--accent2);
                    font-weight: 700;
                }}
                .grid {{
                    display: grid;
                    grid-template-columns: 1fr;
                    gap: 14px;
                }}
                .plot {{
                    background: var(--card);
                    border-radius: 12px;
                    padding: 8px;
                    box-shadow: 0 6px 18px rgba(0, 0, 0, 0.08);
                }}
            </style>
        </head>
        <body>
            <div class=\"container\">
                <h1>Dashboard Interactivo con Filtros</h1>
                <p class=\"subtitle\">Filtra por campana, canal y tipo de cliente para actualizar KPIs y graficos en tiempo real.</p>

                <div class=\"filters\">
                    <div class=\"filter-box\">
                        <label for=\"campaignFilter\">Campana</label>
                        <select id=\"campaignFilter\"></select>
                    </div>
                    <div class=\"filter-box\">
                        <label for=\"channelFilter\">Canal</label>
                        <select id=\"channelFilter\"></select>
                    </div>
                    <div class=\"filter-box\">
                        <label for=\"customerFilter\">Tipo de cliente</label>
                        <select id=\"customerFilter\"></select>
                    </div>
                </div>

                <div class=\"cards\" id=\"kpiCards\"></div>

                <div class=\"grid\">
                    <div class=\"plot\" id=\"plotRevenueChannel\"></div>
                    <div class=\"plot\" id=\"plotConvCampaign\"></div>
                    <div class=\"plot\" id=\"plotRevenueCombo\"></div>
                </div>
            </div>

            <script>
                const rawData = {records_json};

                const campaignFilter = document.getElementById('campaignFilter');
                const channelFilter = document.getElementById('channelFilter');
                const customerFilter = document.getElementById('customerFilter');

                function uniqueValues(key) {{
                    return [...new Set(rawData.map(r => r[key]))].sort();
                }}

                function fillSelect(select, values) {{
                    select.innerHTML = '';
                    const allOpt = document.createElement('option');
                    allOpt.value = 'Todos';
                    allOpt.textContent = 'Todos';
                    select.appendChild(allOpt);
                    values.forEach(v => {{
                        const opt = document.createElement('option');
                        opt.value = v;
                        opt.textContent = v;
                        select.appendChild(opt);
                    }});
                }}

                function filterData() {{
                    const c = campaignFilter.value;
                    const ch = channelFilter.value;
                    const ct = customerFilter.value;
                    return rawData.filter(r =>
                        (c === 'Todos' || r.campaign === c) &&
                        (ch === 'Todos' || r.channel === ch) &&
                        (ct === 'Todos' || r.customer_type === ct)
                    );
                }}

                function groupSum(data, key, valueKey) {{
                    const m = new Map();
                    data.forEach(r => m.set(r[key], (m.get(r[key]) || 0) + r[valueKey]));
                    return [...m.entries()].map(([k, v]) => ({{k, v}}));
                }}

                function groupMean(data, key, valueKey) {{
                    const m = new Map();
                    data.forEach(r => {{
                        const current = m.get(r[key]) || {{sum: 0, n: 0}};
                        current.sum += r[valueKey];
                        current.n += 1;
                        m.set(r[key], current);
                    }});
                    return [...m.entries()].map(([k, obj]) => ({{k, v: obj.n ? obj.sum / obj.n : 0}}));
                }}

                function render() {{
                    const data = filterData();
                    const interactions = data.length;
                    const totalRevenue = data.reduce((acc, r) => acc + r.sales, 0);
                    const conversionRate = interactions ? data.reduce((acc, r) => acc + r.converted, 0) / interactions : 0;
                    const avgTime = interactions ? data.reduce((acc, r) => acc + r.time_on_site, 0) / interactions : 0;

                    const converted = data.filter(r => r.converted === 1);
                    const avgTicket = converted.length ? converted.reduce((acc, r) => acc + r.sales, 0) / converted.length : 0;

                    document.getElementById('kpiCards').innerHTML = `
                        <div class=\"card\"><h3>Interacciones</h3><p>${{interactions.toLocaleString()}}</p></div>
                        <div class=\"card\"><h3>Ingresos</h3><p>$${{totalRevenue.toLocaleString(undefined, {{maximumFractionDigits: 2}})}}</p></div>
                        <div class=\"card\"><h3>Conversion</h3><p>${{(conversionRate * 100).toFixed(2)}}%</p></div>
                        <div class=\"card\"><h3>Tiempo Promedio</h3><p>${{avgTime.toFixed(1)}}s</p></div>
                        <div class=\"card\"><h3>Ticket Promedio</h3><p>$${{avgTicket.toFixed(2)}}</p></div>
                    `;

                    const revenueByChannel = groupSum(data, 'channel', 'sales').sort((a, b) => b.v - a.v);
                    const convByCampaign = groupMean(data, 'campaign', 'converted').sort((a, b) => b.v - a.v);

                    const revenueByComboMap = new Map();
                    data.forEach(r => {{
                        const combo = `${{r.campaign}} - ${{r.channel}}`;
                        revenueByComboMap.set(combo, (revenueByComboMap.get(combo) || 0) + r.sales);
                    }});
                    const revenueByCombo = [...revenueByComboMap.entries()].map(([k, v]) => ({{k, v}})).sort((a, b) => b.v - a.v);

                    Plotly.react('plotRevenueChannel', [{{
                        type: 'bar',
                        x: revenueByChannel.map(d => d.k),
                        y: revenueByChannel.map(d => d.v),
                        marker: {{color: '#457b9d'}}
                    }}], {{title: 'Ingresos por Canal', yaxis: {{title: 'Ingresos ($)'}}}}, {{responsive: true}});

                    Plotly.react('plotConvCampaign', [{{
                        type: 'bar',
                        x: convByCampaign.map(d => d.k),
                        y: convByCampaign.map(d => d.v * 100),
                        marker: {{color: '#e76f51'}}
                    }}], {{title: 'Conversion por Campana (%)', yaxis: {{title: 'Conversion (%)'}}}}, {{responsive: true}});

                    Plotly.react('plotRevenueCombo', [{{
                        type: 'bar',
                        x: revenueByCombo.map(d => d.k),
                        y: revenueByCombo.map(d => d.v),
                        marker: {{color: '#2a9d8f'}}
                    }}], {{title: 'Ingresos por Combinacion Campana-Canal', yaxis: {{title: 'Ingresos ($)'}}}}, {{responsive: true}});
                }}

                fillSelect(campaignFilter, uniqueValues('campaign'));
                fillSelect(channelFilter, uniqueValues('channel'));
                fillSelect(customerFilter, uniqueValues('customer_type'));

                campaignFilter.addEventListener('change', render);
                channelFilter.addEventListener('change', render);
                customerFilter.addEventListener('change', render);

                render();
            </script>
        </body>
        </html>
        """

        output_path.write_text(html, encoding="utf-8")
        return output_path


def recommendation_bullets(tables: Dict[str, pd.DataFrame]) -> List[str]:
    top_revenue = tables["by_combo"].iloc[0]
    top_new = tables["new_customer_combo"].iloc[0]
    top_channel = tables["by_channel"].iloc[0]

    bullets = [
        (
            f"Escalar inversion en {top_revenue['campaign']} + {top_revenue['channel']} porque es la "
            f"combinacion con mayor ingreso total (${top_revenue['revenue']:,.0f})."
        ),
        (
            f"Priorizar adquisicion en {top_new['campaign']} + {top_new['channel']} donde el ingreso "
            f"de clientes nuevos es mas alto (${top_new['revenue']:,.0f})."
        ),
        (
            f"Mantener {top_channel['channel']} como canal principal de ingreso y optimizar pruebas creativas "
            "para elevar conversion sin perder escala."
        ),
        "Implementar landing pages y mensajes por canal diferenciando clientes nuevos y existentes.",
        "Monitorear KPIs semanalmente (conversion, ingreso y ticket) para reasignar presupuesto en forma dinamica.",
    ]
    return bullets


def build_presentation(
    kpis: Dict[str, float],
    tables: Dict[str, pd.DataFrame],
    chart_paths: Dict[str, Path],
) -> Path:
    output_path = OUTPUT_DIR / "Marketing_Campaign_Executive_Presentation.pptx"
    prs = Presentation()

    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "Rendimiento de Campanas de Marketing"
    slide.placeholders[1].text = "Brief ejecutivo: insights por campana, canal y tipo de cliente"

    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Resumen de KPIs"
    body = slide.shapes.placeholders[1].text_frame
    body.text = f"Interacciones totales: {int(kpis['interactions']):,}"
    for line in [
        f"Ingresos totales: ${kpis['total_revenue']:,.2f}",
        f"Tasa de conversion: {kpis['conversion_rate'] * 100:.2f}%",
        f"Tiempo promedio en sitio: {kpis['avg_time_on_site']:.1f} seg",
        f"Ticket promedio (convertidos): ${kpis['avg_sale_converted']:.2f}",
    ]:
        p = body.add_paragraph()
        p.text = line

    for title, key in [
        ("Ingresos por Canal y Campana", "revenue_by_channel_campaign"),
        ("Mapa de Calor de Conversion", "conversion_heatmap"),
        ("Ingresos por Tipo de Cliente", "customer_revenue"),
    ]:
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = title
        slide.shapes.add_picture(str(chart_paths[key]), PptInches(0.6), PptInches(1.2), width=PptInches(12.1))

    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Conclusiones"
    body = slide.shapes.placeholders[1].text_frame
    body.text = "La Campana B lidera en ingreso total debido a mayor escala."
    for line in [
        "La eficiencia de conversion de Campana A es similar, por lo que importa mas la mezcla de canales.",
        "Email es el mayor canal de ingreso; Instagram es el canal mas eficiente en conversion.",
        "El principal impulsor de crecimiento en clientes nuevos es Campana B + Email.",
    ]:
        p = body.add_paragraph()
        p.text = line

    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Recomendaciones"
    body = slide.shapes.placeholders[1].text_frame
    recommendations = recommendation_bullets(tables)
    body.text = recommendations[0]
    for rec in recommendations[1:]:
        p = body.add_paragraph()
        p.text = rec

    prs.save(output_path)
    return output_path


def build_word_report(
    kpis: Dict[str, float],
    tables: Dict[str, pd.DataFrame],
    chart_paths: Dict[str, Path],
) -> Path:
    output_path = OUTPUT_DIR / "Marketing_Campaign_Executive_Report.docx"
    doc = Document()

    doc.add_heading("Informe Ejecutivo de Campanas de Marketing", level=0)
    doc.add_paragraph(
        "Este informe resume los resultados del analisis exploratorio de datos para dos campanas "
        "de marketing en canales Email, Instagram y Web Banner."
    )

    doc.add_heading("1. Resumen de KPIs", level=1)
    kpi_table = doc.add_table(rows=1, cols=2)
    kpi_table.style = "Light List"
    hdr_cells = kpi_table.rows[0].cells
    hdr_cells[0].text = "Metrica"
    hdr_cells[1].text = "Valor"

    kpi_items = [
        ("Interacciones totales", f"{int(kpis['interactions']):,}"),
        ("Ingresos totales", f"${kpis['total_revenue']:,.2f}"),
        ("Tasa de conversion", f"{kpis['conversion_rate'] * 100:.2f}%"),
        ("Tiempo promedio en sitio", f"{kpis['avg_time_on_site']:.2f} seg"),
        ("Ticket promedio (convertidos)", f"${kpis['avg_sale_converted']:.2f}"),
    ]

    for metric, value in kpi_items:
        row_cells = kpi_table.add_row().cells
        row_cells[0].text = metric
        row_cells[1].text = value

    doc.add_heading("2. Analisis Visual", level=1)
    for chart_name in [
        "revenue_by_channel_campaign",
        "conversion_heatmap",
        "customer_revenue",
        "sales_distribution",
        "time_on_site_violin",
    ]:
        doc.add_paragraph(chart_name.replace("_", " ").title())
        doc.add_picture(str(chart_paths[chart_name]), width=Inches(6.5))

    doc.add_heading("3. Hallazgos Clave", level=1)
    findings = [
        "La Campana B genero el mayor ingreso total con una conversion muy similar a Campana A.",
        "Email produjo la mayor contribucion de ingreso, mientras Instagram logro la mayor eficiencia de conversion.",
        "El ingreso por conversion es estable; el crecimiento depende de aumentar trafico calificado y mejorar conversion.",
        "El crecimiento en clientes nuevos es mas fuerte en Campana B + Email.",
    ]
    for f in findings:
        doc.add_paragraph(f, style="List Bullet")

    doc.add_heading("4. Recomendaciones", level=1)
    for rec in recommendation_bullets(tables):
        doc.add_paragraph(rec, style="List Number")

    doc.add_heading("5. Plan de Accion Inmediato", level=1)
    for step in [
        "Reasignar parte del presupuesto hacia Campana B + Email y validar el impacto con una prueba controlada de 4 semanas.",
        "Ejecutar pruebas A/B en landing pages por canal con mensajes separados para clientes nuevos y existentes.",
        "Establecer una revision semanal de desempeno con umbrales de KPI y alertas automaticas.",
    ]:
        doc.add_paragraph(step, style="List Bullet")

    doc.save(output_path)
    return output_path


def build_pdf_report(
    kpis: Dict[str, float],
    tables: Dict[str, pd.DataFrame],
    chart_paths: Dict[str, Path],
) -> Path:
    output_path = OUTPUT_DIR / "Marketing_Campaign_Executive_Report.pdf"

    doc = SimpleDocTemplate(
        str(output_path),
        pagesize=A4,
        rightMargin=1.6 * cm,
        leftMargin=1.6 * cm,
        topMargin=1.6 * cm,
        bottomMargin=1.6 * cm,
    )
    styles = getSampleStyleSheet()
    story = []

    story.append(Paragraph("Informe Ejecutivo de Campanas de Marketing", styles["Title"]))
    story.append(Spacer(1, 0.4 * cm))
    story.append(
        Paragraph(
            "Este informe presenta hallazgos del analisis exploratorio para dos campanas "
            "de marketing y sus canales de difusion.",
            styles["BodyText"],
        )
    )
    story.append(Spacer(1, 0.5 * cm))

    story.append(Paragraph("1. KPIs Ejecutivos", styles["Heading2"]))
    kpi_data = [
        ["Metrica", "Valor"],
        ["Interacciones totales", f"{int(kpis['interactions']):,}"],
        ["Ingresos totales", f"${kpis['total_revenue']:,.2f}"],
        ["Tasa de conversion", f"{kpis['conversion_rate'] * 100:.2f}%"],
        ["Tiempo promedio en sitio", f"{kpis['avg_time_on_site']:.2f} seg"],
        ["Ticket promedio (convertidos)", f"${kpis['avg_sale_converted']:.2f}"],
    ]
    table = Table(kpi_data, colWidths=[8.0 * cm, 7.0 * cm])
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#457b9d")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("BACKGROUND", (0, 1), (-1, -1), colors.whitesmoke),
            ]
        )
    )
    story.append(table)
    story.append(Spacer(1, 0.5 * cm))

    story.append(Paragraph("2. Hallazgos Clave", styles["Heading2"]))
    findings = [
        "La Campana B lidera en ingreso total por mayor escala de interacciones.",
        "Email concentra el mayor ingreso total y Instagram muestra la mayor eficiencia de conversion.",
        "Campana B + Email es la combinacion de mayor impacto para clientes nuevos.",
        "El ticket promedio es estable; la oportunidad principal esta en trafico calificado y conversion.",
    ]
    for finding in findings:
        story.append(Paragraph(f"- {finding}", styles["BodyText"]))
    story.append(Spacer(1, 0.5 * cm))

    story.append(Paragraph("3. Recomendaciones", styles["Heading2"]))
    for rec in recommendation_bullets(tables):
        story.append(Paragraph(f"- {rec}", styles["BodyText"]))
    story.append(Spacer(1, 0.5 * cm))

    story.append(Paragraph("4. Graficos", styles["Heading2"]))
    for chart_key in ["revenue_by_channel_campaign", "conversion_heatmap", "customer_revenue"]:
        story.append(Paragraph(chart_key.replace("_", " ").title(), styles["Heading3"]))
        story.append(PdfImage(str(chart_paths[chart_key]), width=16.5 * cm, height=9.3 * cm))
        story.append(Spacer(1, 0.3 * cm))

    doc.build(story)
    return output_path


def build_markdown_summary(kpis: Dict[str, float], tables: Dict[str, pd.DataFrame]) -> Path:
    output_path = OUTPUT_DIR / "marketing_campaign_analysis_summary.md"

    top_combo = tables["by_combo"].iloc[0]
    top_new_combo = tables["new_customer_combo"].iloc[0]

    text = f"""
# Resumen de Analisis de Campanas de Marketing

## Dataset
- Filas: {int(kpis['interactions']):,}
- Campanas: A y B
- Canales: Email, Instagram, Web Banner

## KPIs Ejecutivos
- Ingresos totales: ${kpis['total_revenue']:,.2f}
- Tasa de conversion: {kpis['conversion_rate'] * 100:.2f}%
- Tiempo promedio en sitio: {kpis['avg_time_on_site']:.2f} segundos
- Ticket promedio (solo convertidos): ${kpis['avg_sale_converted']:.2f}

## Insights de Rendimiento
- Mayor combinacion de ingreso total: {top_combo['campaign']} + {top_combo['channel']} (${top_combo['revenue']:,.2f})
- Mayor combinacion de ingreso en clientes nuevos: {top_new_combo['campaign']} + {top_new_combo['channel']} (${top_new_combo['revenue']:,.2f})
- Campana B es la mayor contribuyente de ingreso por escala.
- Email es el canal lider en ingresos; Instagram presenta la mayor eficiencia de conversion.

## Recomendaciones
1. Aumentar la inversion en Campana B + Email como motor principal de crecimiento.
2. Mantener Email como canal de escala y mejorar conversion con pruebas creativas segmentadas.
3. Disenar journeys por canal diferenciando clientes nuevos y existentes.
4. Implementar gobierno semanal de KPIs para reasignar presupuesto rapidamente.
""".strip()

    output_path.write_text(text + "\n", encoding="utf-8")
    return output_path


def build_notebook() -> Path:
    output_path = OUTPUT_DIR / "marketing_campaign_analysis.ipynb"
    nb = nbf.v4.new_notebook()

    markdown_intro = """
# Notebook de Analisis de Campanas de Marketing

Este notebook documenta el flujo completo:
1. Carga de datos y chequeos de calidad.
2. Analisis exploratorio y calculo de KPIs.
3. Visualizaciones para storytelling ejecutivo.
4. Exportacion de dashboard, dashboard interactivo, presentacion e informe (DOCX/PDF).
""".strip()

    code_setup = """
from pathlib import Path
import pandas as pd
import matplotlib.pyplot as plt
import seaborn as sns

base_dir = Path.cwd()
if (base_dir / 'Marketing_Campaign_Data.csv').exists():
    data_path = base_dir / 'Marketing_Campaign_Data.csv'
else:
    data_path = base_dir / 'Marketing Analysis' / 'Marketing_Campaign_Data.csv'

df = pd.read_csv(data_path)
df = df.rename(columns={
    'Interaction ID': 'interaction_id',
    'Campaign Type': 'campaign',
    'Channel': 'channel',
    'Customer Type': 'customer_type',
    'Converted (1=yes, 0=no)': 'converted',
    'Time on Site (seconds)': 'time_on_site',
    'Sales ($)': 'sales'
})

sns.set_theme(style='whitegrid')
df.head()
""".strip()

    code_quality = """
print('Shape:', df.shape)
print('\nNull values:')
print(df.isna().sum())
print('\nDtypes:')
print(df.dtypes)
""".strip()

    code_kpis = """
kpis = {
    'interactions': len(df),
    'total_revenue': df['sales'].sum(),
    'conversion_rate': df['converted'].mean(),
    'avg_time_on_site': df['time_on_site'].mean(),
    'avg_ticket_converted': df.loc[df['converted'] == 1, 'sales'].mean()
}

kpis
""".strip()

    code_grouping = """
by_campaign = df.groupby('campaign', as_index=False).agg(
    interactions=('interaction_id', 'count'),
    conversion_rate=('converted', 'mean'),
    revenue=('sales', 'sum')
).sort_values('revenue', ascending=False)

by_channel = df.groupby('channel', as_index=False).agg(
    interactions=('interaction_id', 'count'),
    conversion_rate=('converted', 'mean'),
    revenue=('sales', 'sum')
).sort_values('revenue', ascending=False)

by_combo = df.groupby(['campaign', 'channel'], as_index=False).agg(
    interactions=('interaction_id', 'count'),
    conversion_rate=('converted', 'mean'),
    revenue=('sales', 'sum')
).sort_values('revenue', ascending=False)

by_campaign, by_channel, by_combo.head()
""".strip()

    code_plot_1 = """
plt.figure(figsize=(10,6))
sns.barplot(data=by_channel, x='channel', y='revenue', palette='Set2')
plt.title('Revenue by Channel')
plt.ylabel('Revenue ($)')
plt.xlabel('Channel')
plt.show()
""".strip()

    code_plot_2 = """
pivot_conv = by_combo.pivot(index='channel', columns='campaign', values='conversion_rate') * 100
plt.figure(figsize=(8,5))
sns.heatmap(pivot_conv, annot=True, fmt='.2f', cmap='YlGnBu')
plt.title('Conversion Rate (%) by Campaign and Channel')
plt.show()
""".strip()

    code_conclusions = """
print('Conclusiones principales:')
print('- La Campana B genera mas ingreso total por mayor volumen de interacciones.')
print('- Email es el canal con mayor ingreso total.')
print('- Campana B + Email es la combinacion mas fuerte para crecimiento de clientes nuevos.')
""".strip()

    code_export = """
import subprocess
import sys

# Ejecuta este paso para regenerar dashboard, dashboard interactivo, pptx, docx y pdf desde el notebook.
subprocess.run([sys.executable, 'generate_marketing_deliverables.py'], check=False)
""".strip()

    nb.cells = [
        nbf.v4.new_markdown_cell(markdown_intro),
        nbf.v4.new_code_cell(code_setup),
        nbf.v4.new_code_cell(code_quality),
        nbf.v4.new_code_cell(code_kpis),
        nbf.v4.new_code_cell(code_grouping),
        nbf.v4.new_code_cell(code_plot_1),
        nbf.v4.new_code_cell(code_plot_2),
        nbf.v4.new_code_cell(code_conclusions),
        nbf.v4.new_code_cell(code_export),
    ]

    nbf.write(nb, output_path)
    return output_path


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    CHART_DIR.mkdir(parents=True, exist_ok=True)

    df = load_data()
    kpis = compute_kpis(df)
    tables = grouped_tables(df)
    chart_paths = save_charts(df, tables)

    dashboard_path = build_dashboard_html(df, tables, kpis)
    filtered_dashboard_path = build_interactive_filtered_dashboard(df)
    ppt_path = build_presentation(kpis, tables, chart_paths)
    docx_path = build_word_report(kpis, tables, chart_paths)
    pdf_path = build_pdf_report(kpis, tables, chart_paths)
    md_path = build_markdown_summary(kpis, tables)
    notebook_path = build_notebook()

    print("Generated deliverables:")
    for p in [
        dashboard_path,
        filtered_dashboard_path,
        ppt_path,
        docx_path,
        pdf_path,
        md_path,
        notebook_path,
    ]:
        print(f"- {p}")


if __name__ == "__main__":
    main()
